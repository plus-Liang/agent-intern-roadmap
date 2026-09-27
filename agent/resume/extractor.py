# -*- coding: utf-8 -*-
"""简历附件文字提取（Chainlit 上传的 PDF / Word）。

放在 `agent/resume/` 而不是 `agent/app.py`，是为了让纯函数能离线单测
（`app.py` 有 import 副作用：建库、读密钥、注册 Chainlit 回调）。

约定：
- `extract_text(path, name=None)` → `(text, error)`，**从不抛异常**；
  `text` 非空表示成功，否则 `error` 是给用户看的一句话原因。
- 只认 PDF / Word（.docx）；`.doc`（老二进制格式）python-docx 打不开 → 归到不支持。
- 图片（.png/.jpg/...）单独由 `is_image()` 识别，调用方要给"模型不支持图片"的明确回复。
"""
from __future__ import annotations

import zipfile
from pathlib import Path
from typing import Optional, Tuple

# 可解析的扩展名（解析器惰性 import，PDF 用户不装 docx 也能聊别的）
PDF_EXTS = {".pdf"}
WORD_EXTS = {".docx"}

# 明确走"模型不支持图片"话术的扩展名
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tif", ".tiff"}

SUPPORTED_HINT = "目前只支持 **PDF（.pdf）** 和 **Word（.docx）**"


def _suffix(path: Path, name: Optional[str] = None) -> str:
    """按 扩展名 → 文件名 → 文件头 的顺序判断格式。

    Chainlit 落盘的临时文件是 `uuid.<ext>`，扩展名通常都在；
    但用户也可能上传没有扩展名的 PDF，所以用 name / 文件头兜底。
    """
    for candidate in (path.name, name or ""):
        if "." in candidate:
            suffix = Path(candidate).suffix.lower()
            if suffix:
                return suffix
    try:
        with open(path, "rb") as fh:
            head = fh.read(8)
        if head.startswith(b"%PDF"):
            return ".pdf"
        if head.startswith(b"PK\x03\x04"):
            with zipfile.ZipFile(path) as zf:
                names = zf.namelist()
            if any(n.startswith("word/") for n in names):
                return ".docx"
    except Exception:                                 # noqa: BLE001 - 只是猜格式
        return ""
    return ""


def is_image(path, name: Optional[str] = None) -> bool:
    """这张附件是不是图片（图片一律不解析，走"模型不支持图片"话术）。"""
    p = Path(path)
    if p.exists():
        suffix = _suffix(p, name)
    else:
        raw = name or p.name
        suffix = Path(raw).suffix.lower() if "." in raw else ""
    return suffix in IMAGE_EXTS


def _extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n".join((page.extract_text() or "") for page in reader.pages).strip()


def _extract_docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    lines = [p.text for p in document.paragraphs]
    # 简历里大量信息在表格里（教育 / 技能常排版成表），段落之外要把表格也抓上
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            line = " | ".join([c for c in cells if c])
            if line:
                lines.append(line)
    return "\n".join(lines).strip()


def extract_text(path, name: Optional[str] = None) -> Tuple[str, Optional[str]]:
    """从 PDF / Word 附件提取纯文本。

    返回 `(text, error)`：成功 `(text, None)`，失败 `("", 给用户看的原因)`。
    **任何异常都被吞成 error 字符串**，调用方只看返回值。
    """
    p = Path(path)
    if not p.exists():
        return "", f"附件读取失败（文件不存在：{p.name}）"

    try:
        suffix = _suffix(p, name)
    except Exception as e:                            # noqa: BLE001
        return "", f"附件读取失败（{type(e).__name__}: {e}）"

    try:
        if suffix in PDF_EXTS:
            text = _extract_pdf(p)
        elif suffix in WORD_EXTS:
            text = _extract_docx(p)
        elif suffix == ".doc":
            return "", (
                "暂不支持老版 `.doc`（请另存为 `.docx` 再上传）；"
                f"{SUPPORTED_HINT}，也可以直接粘贴简历文本。"
            )
        else:
            shown = suffix or "未知格式"
            return "", (
                f"暂不支持 `{shown}` 格式（{SUPPORTED_HINT}）；"
                "也可以直接粘贴简历文本。"
            )
    except Exception as e:                            # noqa: BLE001 - 解析失败降级为提示
        return "", (
            f"附件解析失败（{type(e).__name__}），可能是加密 / 扫描件 / 文件损坏；"
            "请粘贴简历文本，或换一份文件重试。"
        )

    if not text or not text.strip():
        return "", (
            "附件里没有提取到可读文字（扫描件 / 纯图片 PDF 提取不出文字）；"
            "请粘贴简历文本。"
        )
    return text.strip(), None
